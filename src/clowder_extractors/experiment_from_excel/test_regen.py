import os

import pytest
from openpyxl import Workbook

from remat_experiment_from_excel import (
    PARENT_URL_HEADER,
    excel_to_json,
    extract_regen,
    parse_dataset_id,
    read_decon_from_worksheet,
    read_oligomers_from_worksheet,
    read_postproc_from_worksheet,
    read_results_from_worksheet,
)

# Repo root, four levels up from this test file
# (.../src/clowder_extractors/experiment_from_excel/test_regen.py).
#  cd /Users/tks/MS-UIUC/RA/Code/clowder-extractors/src/clowder_extractors/experiment_from_excel ../../../venv/bin/python -m pytest test_regen.py -v
REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
REGEN_DIR = os.path.join(REPO_ROOT, "Regen-data")

GEN0 = os.path.join(REGEN_DIR, "Gen 0", "Data Entry_Gen0.xlsx")
GEN1 = os.path.join(REGEN_DIR, "Gen 1", "Data Entry_Gen1.xlsx")
GEN2 = os.path.join(REGEN_DIR, "Gen 2", "Data Entry_Gen2.xlsx")
NON_REGEN = os.path.join(REPO_ROOT, "DRD1-81-12 Data Entry.xlsx")

GEN0_URL = (
    "https://re-mat.clowderframework.org/datasets/"
    "6a3e24a1e4b061139149da1f?space=6a3e1dbde4b061139149d7b3"
)
GEN1_URL = (
    "https://re-mat.clowderframework.org/datasets/"
    "6a3e24a2e4b061139149da35?space=6a3e1dbde4b061139149d7b3"
)

has_samples = pytest.mark.skipif(
    not os.path.isdir(REGEN_DIR), reason="Regen-data sample files not present"
)


# --- parse_dataset_id ---------------------------------------------------------


def test_parse_dataset_id_with_space_query():
    assert parse_dataset_id(GEN0_URL) == "6a3e24a1e4b061139149da1f"


def test_parse_dataset_id_without_query():
    url = "https://re-mat.clowderframework.org/datasets/abc123"
    assert parse_dataset_id(url) == "abc123"


def test_parse_dataset_id_malformed():
    assert parse_dataset_id("https://example.com/foo/bar") is None


def test_parse_dataset_id_empty():
    assert parse_dataset_id(None) is None
    assert parse_dataset_id("") is None


# --- read_oligomers_from_worksheet -------------------------------------------


def _make_oligomers_ws(header, parent_rows, procedure_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "oligomers"
    ws.append(header)
    for row in parent_rows:
        ws.append(row)
    ws.append(["PROCEDURE"])
    for row in procedure_rows:
        ws.append(row)
    return ws


def test_oligomers_root():
    ws = _make_oligomers_ws(
        ["Oligo ID", "Measured mass (g)"],
        [[None, None]],
        [["Prepared in GloveBox?", "NO"], ["Mixing type", "N/A"]],
    )
    result = read_oligomers_from_worksheet(ws)
    assert result["is_root"] is True
    assert result["parents"] == []
    assert result["procedure"]["Prepared in GloveBox?"] == "NO"
    assert result["procedure"]["Mixing type"] == "N/A"


def test_oligomers_child_with_url():
    ws = _make_oligomers_ws(
        ["Oligo ID", "Measured mass (g)", PARENT_URL_HEADER],
        [["Gen 0", 2.027, GEN0_URL]],
        [["Prepared in GloveBox?", "NO"]],
    )
    result = read_oligomers_from_worksheet(ws)
    assert result["is_root"] is False
    assert result["parents"] == [
        {
            "oligo_id": "Gen 0",
            "dataset_id": "6a3e24a1e4b061139149da1f",
            "dataset_url": GEN0_URL,
            "Measured mass (g)": 2.027,
        }
    ]
    assert result["procedure"]["Prepared in GloveBox?"] == "NO"


def test_oligomers_child_pre_issue19_no_column_c():
    # Parents listed in column A but no Parent Dataset URL column yet.
    ws = _make_oligomers_ws(
        ["Oligo ID", "Measured mass (g)"],
        [["Gen 0", 2.027]],
        [],
    )
    result = read_oligomers_from_worksheet(ws)
    assert result["is_root"] is False
    assert result["parents"] == [
        {
            "oligo_id": "Gen 0",
            "dataset_id": None,
            "dataset_url": None,
            "Measured mass (g)": 2.027,
        }
    ]


def test_oligomers_multiple_parents():
    ws = _make_oligomers_ws(
        ["Oligo ID", "Measured mass (g)", PARENT_URL_HEADER],
        [
            ["Gen 0", 1.5, GEN0_URL],
            ["Gen 1", 2.0, GEN1_URL],
        ],
        [["Prepared in GloveBox?", "YES"]],
    )
    result = read_oligomers_from_worksheet(ws)
    assert result["is_root"] is False
    assert len(result["parents"]) == 2
    assert result["parents"][0]["oligo_id"] == "Gen 0"
    assert result["parents"][0]["dataset_id"] == "6a3e24a1e4b061139149da1f"
    assert result["parents"][0]["Measured mass (g)"] == 1.5
    assert result["parents"][1]["oligo_id"] == "Gen 1"
    assert result["parents"][1]["dataset_id"] == "6a3e24a2e4b061139149da35"
    assert result["parents"][1]["Measured mass (g)"] == 2.0


def test_oligomers_extra_weight_column_omitted_when_null():
    # If weight is None for a row it should not appear in that parent dict.
    ws = _make_oligomers_ws(
        ["Oligo ID", "Measured mass (g)", PARENT_URL_HEADER],
        [["Gen 0", None, GEN0_URL]],
        [],
    )
    result = read_oligomers_from_worksheet(ws)
    assert "Measured mass (g)" not in result["parents"][0]


def test_oligomers_stops_at_procedure():
    ws = _make_oligomers_ws(
        ["Oligo ID", "Measured mass (g)", PARENT_URL_HEADER],
        [["Gen 0", 2.027, GEN0_URL]],
        [["Mixing time (mins)", 5]],
    )
    result = read_oligomers_from_worksheet(ws)
    assert len(result["parents"]) == 1
    assert "Mixing time (mins)" in result["procedure"]
    # The procedure key must not leak into parents.
    assert all(p["oligo_id"] != "Mixing time (mins)" for p in result["parents"])


# --- read_decon_from_worksheet -----------------------------------------------


def test_decon_reagents_and_procedure():
    wb = Workbook()
    ws = wb.active
    ws.title = "ReGen_Decon"
    ws.append(["Name", "SMILES", "Measured volume (mL)", "Supplier", "Lot #"])
    ws.append(["TBAF", "[F-].CCCC", 115, None, None])
    ws.append(["PROCEDURE"])
    ws.append(["Prepared in GloveBox?", "NO"])
    ws.append(["Stirr", "NO"])

    result = read_decon_from_worksheet(ws)
    assert result["reagents"] == [
        {
            "Name": "TBAF",
            "SMILES": "[F-].CCCC",
            "Measured volume (mL)": 115,
            "Supplier": None,
            "Lot #": None,
        }
    ]
    assert result["procedure"] == {"Prepared in GloveBox?": "NO", "Stirr": "NO"}


def test_decon_empty_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = "ReGen_Decon"
    assert read_decon_from_worksheet(ws) == {"reagents": [], "procedure": {}}


# --- read_results_from_worksheet ---------------------------------------------


def test_results_skips_valueless_labels():
    wb = Workbook()
    ws = wb.active
    ws.title = "ReGen_Results"
    ws.append(["Oligomer mass (g)", 1.5])
    ws.append(["Inital mass (g)", None])
    ws.append(["CLOWDER SHOULD CALCULATE YIELD MASS RATIOS", None])

    result = read_results_from_worksheet(ws)
    assert result == {"Oligomer mass (g)": 1.5}


# --- read_postproc_from_worksheet --------------------------------------------


def test_postproc_steps_drop_placeholders():
    wb = Workbook()
    ws = wb.active
    ws.title = "ReGen_PostProc"
    ws.append(["STEP 1", None, "STEP 2", None])
    ws.append(["Process", "Precipitate", "Process", "Filter"])
    ws.append(["Temperature (C) ", 20, "Temperature (C) ", 20])
    ws.append(["AntiSol 1", "Methanol", "Mesh size", "N/A"])
    ws.append(["Mixing", "#N/A", "N/A", "#N/A"])

    steps = read_postproc_from_worksheet(ws)
    assert steps == {
        "STEP 1": {
            "Process": "Precipitate",
            "Temperature (C) ": 20,
            "AntiSol 1": "Methanol",
        },
        "STEP 2": {"Process": "Filter", "Temperature (C) ": 20},
    }


def test_postproc_empty_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = "ReGen_PostProc"
    assert read_postproc_from_worksheet(ws) == {}


# --- extract_regen isolation -------------------------------------------------


def test_extract_regen_isolates_failing_parser(monkeypatch):
    import remat_experiment_from_excel as mod

    def boom(ws):
        raise RuntimeError("kaboom")

    wb = Workbook()
    wb.active.title = "oligomers"
    wb.active.append(["Oligo ID", "Measured mass (g)"])
    wb.create_sheet("ReGen_Decon")

    monkeypatch.setattr(
        mod,
        "REGEN_TAB_PARSERS",
        {
            "oligomers": ("oligomers", mod.read_oligomers_from_worksheet),
            "ReGen_Decon": ("decon", boom),
        },
    )

    regen = extract_regen(wb)
    # Failing parser is swallowed; the working one still emits.
    assert "oligomers" in regen
    assert "decon" not in regen


# --- real-file integration ---------------------------------------------------


@has_samples
def test_gen0_is_root():
    result = excel_to_json(GEN0)
    regen = result["regen"]
    assert regen["oligomers"]["is_root"] is True
    assert regen["oligomers"]["parents"] == []
    # All implemented ReGen tabs are extracted as sibling subkeys.
    assert set(regen) == {"oligomers", "decon", "postproc", "results"}
    assert regen["decon"]["reagents"][0]["Name"].startswith("Tetrebutylammonium")
    assert {v["Process"] for v in regen["postproc"].values()} == {
        "Precipitate",
        "Filter",
        "Drying",
    }
    # Standard extraction still runs and merges in.
    assert "Batch ID" in result
    assert "inputs" in result
    assert "procedure" in result


@has_samples
def test_gen1_parent_is_gen0():
    result = excel_to_json(GEN1)
    parents = result["regen"]["oligomers"]["parents"]
    assert len(parents) == 1
    assert parents[0]["oligo_id"] == "Gen 0"
    assert parents[0]["dataset_id"] == "6a3e24a1e4b061139149da1f"
    assert "Batch ID" in result


@has_samples
def test_gen2_parent_is_gen1():
    result = excel_to_json(GEN2)
    parents = result["regen"]["oligomers"]["parents"]
    assert len(parents) == 1
    assert parents[0]["oligo_id"] == "Gen 1"
    assert parents[0]["dataset_id"] == "6a3e24a2e4b061139149da35"


@has_samples
def test_non_regen_has_no_regen_key():
    result = excel_to_json(NON_REGEN)
    assert "regen" not in result
    assert "Batch ID" in result
